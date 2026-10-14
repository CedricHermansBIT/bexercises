"""Read generated data in a separate, network-isolated grading container."""
import base64
import datetime
import io
import json
import math
import sqlite3
import sys
import zipfile
import zlib

MAX_CELLS = 2500
MAX_ROWS = 1000


def typed(value, cell_type=None):
    if value is None:
        return {"type": "null", "value": None}
    if cell_type == "f":
        return {"type": "formula", "value": value}
    if cell_type == "e":
        return {"type": "error", "value": value}
    if isinstance(value, bool):
        return {"type": "boolean", "value": value}
    if isinstance(value, int):
        return {"type": "number", "value": str(value)}
    if isinstance(value, float) and math.isfinite(value):
        return {"type": "number", "value": str(int(value)) if value.is_integer() else repr(value)}
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return {"type": "date", "value": value.isoformat()}
    if isinstance(value, bytes):
        return {"type": "blob", "value": value.hex()}
    if isinstance(value, str):
        return {"type": "string", "value": value}
    raise ValueError("Unsupported or non-finite cell value")


def inspect_xlsx(filename, check):
    import openpyxl
    with zipfile.ZipFile(filename) as archive:
        members = archive.infolist()
        if len(members) > 1000 or sum(item.file_size for item in members) > 16 * 1024 * 1024:
            raise ValueError("Workbook exceeds the expanded archive limit")
    with open(filename, "rb") as source:
        workbook = openpyxl.load_workbook(io.BytesIO(source.read()), data_only=False, keep_links=False)
    try:
        names = sorted(check.get("sheets") or workbook.sheetnames)
        if len(names) > 20:
            raise ValueError("Workbook exceeds 20 checked sheets")
        sheets = []
        previews = []
        count = 0
        for name in names:
            if name not in workbook.sheetnames:
                raise ValueError("Missing worksheet: " + name)
            sheet = workbook[name]
            cells = {}
            # Only populated cells affect content. Formatting on empty cells is ignored.
            # The course image pins openpyxl; sparse storage avoids expanding empty ranges.
            populated = sorted((cell for cell in sheet._cells.values() if cell.value is not None),
                               key=lambda cell: (cell.row, cell.column))
            count += len(populated)
            if count > MAX_CELLS:
                raise ValueError("Workbook exceeds 2500 populated checked cells")
            for cell in populated:
                cells[cell.coordinate] = typed(cell.value, cell.data_type)
            sheets.append({"name": name, "cells": cells})
            previews.append({"name": name, "columns": ["Cell", "Type", "Value"],
                             "rows": [[cell.coordinate, cells[cell.coordinate]["type"],
                                       str(cells[cell.coordinate]["value"])[:120]] for cell in populated[:20]],
                             "totalRows": len(populated)})
        return {"kind": "xlsx", "data": {"sheets": sheets}, "preview": {"tables": previews}}
    finally:
        workbook.close()


def inspect_sqlite(filename, check):
    from pathlib import Path
    connection = sqlite3.connect(Path(filename).resolve().as_uri() + "?mode=ro", uri=True)
    connection.execute("PRAGMA query_only=ON")
    allowed = {sqlite3.SQLITE_SELECT, sqlite3.SQLITE_READ, sqlite3.SQLITE_FUNCTION, sqlite3.SQLITE_RECURSIVE}

    def authorise(action, first, second, _db, _source):
        if action not in allowed or (action == sqlite3.SQLITE_FUNCTION and second == "load_extension"):
            return sqlite3.SQLITE_DENY
        return sqlite3.SQLITE_OK

    connection.set_authorizer(authorise)
    steps = 0

    def progress():
        nonlocal steps
        steps += 1000
        return int(steps > 1000000)

    connection.set_progress_handler(progress, 1000)
    queries = []
    previews = []
    total = 0
    try:
        for query in check["queries"]:
            cursor = connection.execute(query["sql"], query.get("parameters") or [])
            columns = [item[0] for item in cursor.description]
            if len(columns) > 50:
                raise ValueError("SQLite query exceeds 50 columns")
            rows = cursor.fetchmany(MAX_ROWS + 1)
            total += len(rows)
            if total > MAX_ROWS:
                raise ValueError("SQLite checks exceed 1000 returned rows")
            data = [[typed(value) for value in row] for row in rows]
            queries.append({"name": query["name"], "columns": columns, "rows": data})
            previews.append({"name": query["name"], "columns": columns,
                             "rows": [[str(value)[:120] if value is not None else "NULL" for value in row]
                                      for row in rows[:20]], "totalRows": len(rows)})
        return {"kind": "sqlite", "data": {"queries": queries}, "preview": {"tables": previews}}
    finally:
        connection.close()


def inspect_image(filename, check):
    from PIL import Image
    Image.MAX_IMAGE_PIXELS = 1000000
    with Image.open(filename) as image:
        if image.width * image.height > 1000000 or getattr(image, "n_frames", 1) != 1:
            raise ValueError("Image must have one frame and at most 1,000,000 pixels")
        # Compare decoded RGBA pixels, not metadata, file format or compression.
        rgba = image.convert("RGBA")
        if check["kind"] == "plot":
            white = Image.new("RGBA", rgba.size, "white")
            white.alpha_composite(rgba)
            rgba = white
        compressed = zlib.compress(rgba.tobytes())
        if len(compressed) > 512 * 1024:
            raise ValueError("Image pixel snapshot exceeds 512 KiB; use a smaller exercise image")
        thumb = rgba.copy()
        thumb.thumbnail((256, 256))
        stream = io.BytesIO()
        thumb.save(stream, format="PNG", optimize=True)
        if stream.tell() > 64 * 1024:
            raise ValueError("Image thumbnail exceeds 64 KiB; use a smaller exercise image")
        return {"kind": check["kind"], "data": {"width": rgba.width, "height": rgba.height,
                    "rgba": base64.b64encode(compressed).decode("ascii")},
                "preview": {"width": rgba.width, "height": rgba.height,
                    "png": base64.b64encode(stream.getvalue()).decode("ascii")}}



def inspect_docx(filename, check):
    from docx import Document
    from docx.text.paragraph import Paragraph
    from docx.table import Table
    with zipfile.ZipFile(filename) as archive:
        members = archive.infolist()
        if len(members) > 1000 or sum(item.file_size for item in members) > 16 * 1024 * 1024:
            raise ValueError("Word document exceeds the expanded archive limit")
    document = Document(filename)
    blocks = []
    for element in document.element.body:
        if element.tag.endswith("}p"):
            paragraph = Paragraph(element, document)
            # Ignore layout-only blank paragraphs; retain order of actual content.
            if not paragraph.text.strip():
                continue
            block = {"kind": "paragraph", "text": paragraph.text}
            if check.get("headings", True):
                style = paragraph.style.name
                block["heading"] = style if style == "Title" or style.startswith("Heading ") else None
            blocks.append(block)
        elif element.tag.endswith("}tbl"):
            table = Table(element, document)
            blocks.append({"kind": "table", "rows": [[cell.text for cell in row.cells] for row in table.rows]})
        if len(blocks) > 200:
            raise ValueError("Word document exceeds 200 content blocks")
    if len(json.dumps(blocks)) > 128 * 1024:
        raise ValueError("Word document content exceeds 128 KiB")
    preview_rows = []
    for index, block in enumerate(blocks):
        if block["kind"] == "paragraph":
            preview_rows.append([str(index + 1), block.get("heading") or "Paragraph", block["text"][:300]])
        else:
            for row in block["rows"]:
                preview_rows.append([str(index + 1), "Table row", " | ".join(row)[:300]])
    return {"kind": "docx", "data": {"blocks": blocks}, "preview": {"tables": [
        {"name": "Document content (not page layout)", "columns": ["Block", "Kind", "Text"],
         "rows": preview_rows[:20], "totalRows": len(preview_rows)}]}}


def png_preview(image):
    image.thumbnail((300, 300))
    stream = io.BytesIO()
    image.save(stream, format="PNG", optimize=True)
    if stream.tell() > 64 * 1024:
        raise ValueError("Preview exceeds 64 KiB")
    return base64.b64encode(stream.getvalue()).decode("ascii")


def inspect_pdf(filename, check):
    import pypdfium2 as pdfium
    pages = []
    previews = []
    total_chars = 0
    with pdfium.PdfDocument(filename) as document:
        if len(document) < 1 or len(document) > 10:
            raise ValueError("PDF must contain 1–10 pages")
        for index in range(len(document)):
            page = document[index]
            try:
                width, height = page.get_size()
                if min(width, height) <= 0 or max(width, height) > 3000:
                    raise ValueError("PDF page size must be at most 3000 points")
                textpage = page.get_textpage()
                try:
                    count = textpage.count_chars()
                    total_chars += count
                    if total_chars > 10000:
                        raise ValueError("PDF exceeds 10,000 text characters")
                    text = " ".join(textpage.get_text_range().split())
                    data = {"text": text}
                    if check.get("layout", False):
                        data["size"] = [width, height]
                        data["characters"] = []
                        for char_index in range(count):
                            char = textpage.get_text_range(char_index, 1)
                            if char.strip():
                                data["characters"].append({"text": char, "box": list(textpage.get_charbox(char_index))})
                    pages.append(data)
                    if index < 3:
                        bitmap = page.render(scale=min(300 / width, 300 / height, 1))
                        try:
                            previews.append({"name": "Page " + str(index + 1), "png": png_preview(bitmap.to_pil()), "text": text[:1000]})
                        finally:
                            bitmap.close()
                finally:
                    textpage.close()
            finally:
                page.close()
    return {"kind": "pdf", "data": {"pages": pages}, "preview": {"pages": previews, "totalPages": len(pages)}}

def main():
    check = json.loads(sys.argv[1])
    try:
        reader = {"xlsx": inspect_xlsx, "sqlite": inspect_sqlite, "image": inspect_image, "plot": inspect_image,
                  "docx": inspect_docx, "pdf": inspect_pdf}[check["kind"]]
        report = reader("artifact.bin", check)
    except Exception as error:
        report = {"kind": check["kind"], "error": type(error).__name__ + ": " + str(error)}
    encoded = json.dumps(report, ensure_ascii=True, allow_nan=False, separators=(",", ":"))
    if len(encoded.encode("utf-8")) > 1024 * 1024:
        encoded = json.dumps({"kind": check["kind"], "error": "Artifact snapshot exceeds 1 MiB"})
    with open("artifact-report.json", "w", encoding="utf-8") as output:
        output.write(encoded)


if __name__ == "__main__":
    main()
